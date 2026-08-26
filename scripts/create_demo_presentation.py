from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "presentations"
LOGO = ROOT / "apps" / "web" / "public" / "center-estates-logo.png"
HERO = ROOT / "apps" / "web" / "public" / "images" / "hero.jpg"
FACILITY = ROOT / "apps" / "web" / "public" / "images" / "facility.jpg"

NAVY = RGBColor(10, 33, 61)
NAVY_2 = RGBColor(18, 54, 91)
RED = RGBColor(222, 35, 42)
TEAL = RGBColor(0, 145, 138)
GREEN = RGBColor(22, 163, 74)
AMBER = RGBColor(234, 134, 20)
INK = RGBColor(18, 32, 52)
# Secondary copy must remain comfortably legible on projectors and exported
# PDFs. This slate is intentionally darker than the web UI's tertiary text.
MUTED = RGBColor(55, 73, 94)
PALE = RGBColor(242, 247, 251)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(211, 222, 233)


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Liberation Sans", margin=0.04,
             valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    r = p.runs[0]
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def base_slide(prs, section, title, subtitle=None, number=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = PALE
    rect(slide, 0, 0, 13.333, 0.22, NAVY, NAVY, False)
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(0.42), Inches(0.38), height=Inches(0.62))
    add_text(slide, "CeEst ERP", 1.15, 0.4, 2.2, 0.28, 18, NAVY, True)
    add_text(slide, "PROPERTY  •  FACILITY  •  PERFORMANCE", 1.15, 0.69, 3.8, 0.2, 8, MUTED, True)
    add_text(slide, section.upper(), 0.55, 1.12, 3.5, 0.3, 10, RED, True)
    add_text(slide, title, 0.55, 1.42, 12.2, 0.52, 27, INK, True)
    if subtitle:
        add_text(slide, subtitle, 0.55, 1.92, 12.0, 0.38, 13, MUTED)
    if number is not None:
        add_text(slide, f"{number:02d}", 12.35, 0.45, 0.42, 0.3, 10, RED, True, PP_ALIGN.RIGHT)
    return slide


def bullets(slide, items, x, y, w, h, size=17, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.05)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.level = 0; p.space_after = Pt(11)
        p.font.name = "Liberation Sans"; p.font.size = Pt(size); p.font.color.rgb = color
        p.text = "•  " + p.text
    return box


def proof_bar(slide, text):
    rect(slide, 0.55, 6.78, 12.23, 0.42, RGBColor(229, 239, 250), RGBColor(229, 239, 250))
    add_text(slide, "PROOF", 0.78, 6.87, 0.7, 0.2, 9, NAVY, True)
    add_text(slide, text, 1.48, 6.84, 11.0, 0.25, 11, NAVY, True)


def card(slide, x, y, w, h, title, body, accent=RED, icon=None):
    rect(slide, x, y, w, h, WHITE, LINE)
    rect(slide, x, y, 0.08, h, accent, accent, False)
    if icon:
        add_text(slide, icon, x+0.22, y+0.18, 0.4, 0.35, 20, accent, True)
        tx = x+0.65
    else:
        tx = x+0.25
    add_text(slide, title, tx, y+0.16, w-(tx-x)-0.2, 0.34, 15, INK, True)
    add_text(slide, body, x+0.25, y+0.62, w-0.5, h-0.82, 11.5, MUTED,
             valign=MSO_ANCHOR.TOP, margin=0.02)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)

    # Cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    if HERO.exists():
        s.shapes.add_picture(str(HERO), 0, 0, width=prs.slide_width, height=prs.slide_height)
    # A bright editorial cover with a narrow red signature bar; the image is
    # allowed to carry the atmosphere instead of being buried under navy.
    rect(s, 0, 0, 5.45, 7.5, WHITE, WHITE, False)
    rect(s, 5.33, 0, 0.12, 7.5, RED, RED, False)
    if LOGO.exists(): s.shapes.add_picture(str(LOGO), Inches(0.68), Inches(0.55), height=Inches(0.82))
    add_text(s, "CeEst ERP", 1.62, 0.66, 2.8, 0.4, 23, NAVY, True)
    add_text(s, "LIVE DEMONSTRATION", 0.72, 2.0, 3.7, 0.3, 11, RED, True)
    add_text(s, "One operating record.\nComplete estate control.", 0.68, 2.42, 4.15, 1.35, 29, NAVY, True,
             valign=MSO_ANCHOR.TOP)
    add_text(s, "Property management, facilities operations and staff performance—connected from transaction to report.", 0.72, 4.02, 3.95, 0.9, 14, MUTED,
             valign=MSO_ANCHOR.TOP)
    rect(s, 0.72, 5.28, 3.95, 0.58, RED, RED)
    add_text(s, "RECORD  →  ACT  →  VERIFY  →  REPORT", 0.86, 5.44, 3.67, 0.24, 10, WHITE, True, PP_ALIGN.CENTER)
    add_text(s, "Live demonstration  •  August 2026", 0.72, 6.72, 3.8, 0.25, 10, MUTED)

    s = base_slide(prs, "Demo story", "From property record to management evidence", "A focused 25–30 minute journey—not a tour of every menu.", 1)
    steps = [("01", "RECORD", "Property, tenant, asset and service masters"), ("02", "PLAN", "Billing schedules, PPM and compliance calendars"), ("03", "EXECUTE", "Tasks, tickets, work orders and collections"), ("04", "VERIFY", "Approvals, evidence, SLA and audit trail"), ("05", "MEASURE", "KPIs, trends, ageing and performance"), ("06", "REPORT", "Tenant, manager and landlord outputs")]
    for i,(n,t,b) in enumerate(steps):
        x=0.65+(i%3)*4.15; y=2.65+(i//3)*1.65
        card(s,x,y,3.8,1.3,f"{n}  {t}",b,[RED,TEAL,AMBER,NAVY_2,GREEN,RED][i])
    proof_bar(s, "Every dashboard number is traceable to an underlying transaction and responsible person.")

    s = base_slide(prs, "Command centre", "Start with today’s operating priorities", "Companion gives each manager one place to see what needs attention now.", 2)
    card(s,0.65,2.55,2.9,1.55,"Due today","Planned maintenance, inspections, lease events and assigned tasks.",RED,"◷")
    card(s,3.75,2.55,2.9,1.55,"Needs attention","Overdue actions, SLA risks, expiring compliance and approvals.",AMBER,"!")
    card(s,6.85,2.55,2.9,1.55,"Team delivery","Assignments, progress, evidence and staff performance.",TEAL,"✓")
    card(s,9.95,2.55,2.75,1.55,"Quartermaster AI","Contextual assistance, summaries and operational guidance.",NAVY_2,"✦")
    bullets(s,["Open with the manager’s current property already selected.","Use one alert as the bridge into the live workflow.","Return here at the end to show the action is now reflected."],0.85,4.55,11.8,1.65,16)
    proof_bar(s,"The ERP behaves as an operating command centre, not a passive database.")

    s = base_slide(prs,"Property foundation","A complete, property-scoped operating record","Property → buildings/areas → units → tenants → services → assets.",3)
    card(s,0.65,2.55,3.75,2.7,"Property & buildings","Identity, address, ownership, building hierarchy, lettable area, utilities, documents and contacts.",NAVY_2,"▦")
    card(s,4.58,2.55,3.75,2.7,"Units & billing schedule","Unit use, floor, area, parking, recurring charges, deposits and utility applicability.",RED,"⌂")
    card(s,8.51,2.55,3.75,2.7,"Tenant & lease","Application, contacts, assigned unit, commencement, expiry, escalation and approvals.",TEAL,"◎")
    add_text(s,"All module tabs inherit the selected property—no portfolio totals inside a property workspace.",0.8,5.63,11.7,0.55,16,NAVY,True,PP_ALIGN.CENTER)
    proof_bar(s,"One trusted structure eliminates duplicate setup and enables automatic downstream billing and reporting.")

    s = base_slide(prs,"Property management","Tenancy and billing are configured once","The unit defines what can be billed; the lease defines when rent becomes due.",4)
    card(s,0.65,2.52,3.8,2.8,"1  Configure the unit","Rent basis, service charge, utilities, parking and deposits belong to the unit billing schedule.",NAVY_2)
    card(s,4.77,2.52,3.8,2.8,"2  Create the tenancy","Select the unit, capture lease dates and escalation, then inherit applicable charges automatically.",RED)
    card(s,8.89,2.52,3.8,2.8,"3  Generate invoices","The engine checks readiness and produces separate rent and utility invoices—payable together.",TEAL)
    add_text(s,"Demo action: open one active tenant and explain the source of each charge before opening invoicing.",0.8,5.68,11.7,0.45,15,INK,True,PP_ALIGN.CENTER)
    proof_bar(s,"Billing is driven by tenancy and unit data—not rebuilt manually at month-end.")

    s = base_slide(prs,"Collections","From invoice to receipt, allocation and arrears","A clear financial trail for manager, landlord and tenant.",5)
    flow=[("INVOICE",NAVY_2),("PAYMENT",TEAL),("RECEIPT",GREEN),("ALLOCATION",AMBER),("ARREARS",RED),("REPORT",NAVY)]
    for i,(t,c) in enumerate(flow):
        x=0.65+i*2.05; rect(s,x,2.8,1.72,0.75,c,c); add_text(s,t,x+0.08,3.02,1.56,0.25,11,WHITE,True,PP_ALIGN.CENTER)
        if i<5: add_text(s,"→",x+1.76,3.0,0.26,0.25,18,MUTED,True,PP_ALIGN.CENTER)
    bullets(s,["Show billed, collected and outstanding for the selected property and period.","Open the tenant statement or receipt to prove document quality and traceability.","Explain reminders and recovery actions without overstating unconfigured payment integrations."],0.9,4.25,11.6,1.72,16)
    proof_bar(s,"The manager can explain every balance and the tenant can access the same supporting documents.")

    s = base_slide(prs,"Service charge","Plan, allocate, reconcile and explain","Connect operational requirements to the annual service charge budget.",6)
    card(s,0.7,2.55,2.85,2.45,"Schemes","Define period, allocation basis and applicable units.",NAVY_2,"1")
    card(s,3.72,2.55,2.85,2.45,"Budget lines","Security, cleaning, utilities, repairs and management costs.",RED,"2")
    card(s,6.74,2.55,2.85,2.45,"Allocation","Apportion by area, equal share or configured driver.",TEAL,"3")
    card(s,9.76,2.55,2.85,2.45,"Reconciliation","Compare budget with actual expenditure and explain variance.",AMBER,"4")
    add_text(s,"Use historical averages or projected quantities and rates—then preserve the approval and audit trail.",0.85,5.42,11.6,0.55,15,INK,True,PP_ALIGN.CENTER)
    proof_bar(s,"Service charge planning is linked to real operational requirements and actual expenditure.")

    s = base_slide(prs,"Tenant self-service","A tenant can act—not only view","Invoices, receipts, lease information, notices and maintenance requests in one secure portal.",7)
    card(s,0.65,2.5,3.8,2.75,"Financial self-service","View invoices, balances, receipts and payment history.",NAVY_2,"₭")
    card(s,4.77,2.5,3.8,2.75,"Report an issue","Capture category, priority, description and unit context; receive a traceable reference.",RED,"!")
    card(s,8.89,2.5,3.8,2.75,"Track progress","See open requests, updates, documents and closure evidence.",TEAL,"✓")
    add_text(s,"Live proof: submit a low-risk issue, note the reference, then continue the same record in Facility Management.",0.85,5.6,11.6,0.5,15,INK,True,PP_ALIGN.CENTER)
    proof_bar(s,"A tenant request becomes a property-scoped operational transaction—not an email or informal phone call.")

    s = base_slide(prs,"Facilities operations","Turn a reported defect into controlled work","Ticket → triage → SLA → assignment → evidence → verification → closure.",8)
    labels=["REPORT","TRIAGE","ASSIGN","EXECUTE","VERIFY","CLOSE"]
    colors=[RED,AMBER,NAVY_2,TEAL,GREEN,NAVY]
    for i,(label,c) in enumerate(zip(labels,colors)):
        x=0.7+i*2.05; rect(s,x,2.75,1.68,0.82,c,c); add_text(s,label,x+0.05,3.02,1.58,0.25,11,WHITE,True,PP_ALIGN.CENTER)
        if i<5:add_text(s,"→",x+1.72,3.03,0.28,0.25,18,MUTED,True,PP_ALIGN.CENTER)
    bullets(s,["Time-stamped severity and SLA target", "Responsible staff member or approved service provider", "Photos, notes, costs and verification evidence", "Escalation and notifications when thresholds are at risk"],0.85,4.22,11.7,1.75,16)
    proof_bar(s,"Every maintenance action has ownership, evidence, elapsed time and a measurable outcome.")

    s = base_slide(prs,"Asset lifecycle","Every critical asset has a passport and a plan","From onboarding to preventive maintenance, condition history and replacement decisions.",9)
    if FACILITY.exists():
        s.shapes.add_picture(str(FACILITY), Inches(7.2), Inches(2.38), width=Inches(5.45), height=Inches(3.65))
    bullets(s,["Specification, location, condition, criticality and warranty", "Service history, documents and next planned maintenance", "12-month PPM calendar with due, completed and overdue work", "Inspections, compliance renewals and recurring-failure visibility"],0.75,2.53,5.95,3.2,17)
    proof_bar(s,"The asset register is operational: it drives schedules, work, compliance and lifecycle decisions.")

    s = base_slide(prs,"Service architecture","Define WHAT before assigning WHO","Service Profiles remain stable while providers, contracts and performance can change.",10)
    stages=[("SERVICE\nPROFILE","What is required"),("PROVIDER","Who delivers"),("PROPERTY\nALLOCATION","Where it applies"),("CONTRACT /\nSLA","Expected standard"),("WORK &\nINSPECTION","Evidence"),("KPI / COST","Outcome")]
    for i,(t,b) in enumerate(stages):
        x=0.55+i*2.1; rect(s,x,2.55,1.82,1.9,WHITE,LINE); add_text(s,t,x+0.12,2.78,1.58,0.62,14,[NAVY_2,RED,TEAL,AMBER,NAVY,GREEN][i],True,PP_ALIGN.CENTER); add_text(s,b,x+0.12,3.62,1.58,0.36,10,MUTED,False,PP_ALIGN.CENTER)
        if i<5:add_text(s,"→",x+1.83,3.22,0.25,0.25,17,MUTED,True,PP_ALIGN.CENTER)
    bullets(s,["One profile can use different providers at different properties.","Changing a provider must never erase service requirements or historical performance."],0.95,5.0,11.4,1.05,15)
    proof_bar(s,"Provider selection is separated from the enduring service standard used for SLA and KPI measurement.")

    s = base_slide(prs,"Performance","Transactions become management insight","Use trends and exceptions to answer: What happened? Why? What needs action next?",11)
    card(s,0.65,2.5,2.95,2.75,"PPM compliance","Completed on time, overdue and upcoming preventive work.",GREEN,"✓")
    card(s,3.83,2.5,2.95,2.75,"SLA performance","Response and resolution time by priority, team and provider.",RED,"◷")
    card(s,7.01,2.5,2.95,2.75,"Financial control","Expected vs actual collections, arrears ageing and budget variance.",AMBER,"↗")
    card(s,10.19,2.5,2.5,2.75,"Risk & compliance","Expiries, incidents, recurring defects and approvals.",NAVY_2,"!")
    add_text(s,"Drill from every metric to the source record during the demo.",0.8,5.67,11.8,0.42,15,INK,True,PP_ALIGN.CENTER)
    proof_bar(s,"The system measures performance objectively using the same records created during daily work.")

    s = base_slide(prs,"Reporting","One data source, different audiences","Operational registers, tenant documents and landlord-ready management packs.",12)
    card(s,0.65,2.48,3.75,2.95,"Operational","Work orders, PPM, compliance, incidents, utilities and provider performance.",NAVY_2,"▤")
    card(s,4.58,2.48,3.75,2.95,"Financial","Billing and collection, receipts and expenditure, arrears, VAT and budget versus actual.",RED,"KES")
    card(s,8.51,2.48,3.75,2.95,"Tenancy","Occupancy, vacancy, lease expiry, movements and tenant statements.",TEAL,"⌂")
    bullets(s,["PDF for presentation and audit evidence", "Excel/CSV for analysis and downstream processing", "Property, reporting period, management identity and branding in every document"],0.9,5.62,11.5,0.7,13)
    proof_bar(s,"Management and landlord reports reuse live records instead of separate spreadsheets.")

    s = base_slide(prs,"Live demonstration","Recommended 27-minute run of show","Keep the audience inside one property and one coherent story.",13)
    rows=[("2 min","Companion","Today’s priorities and selected property"),("4 min","Property & tenancy","Unit billing schedule and active lease"),("4 min","Billing & collections","Invoice source, receipt and outstanding balance"),("3 min","Tenant portal","Submit one maintenance issue"),("6 min","Facility workflow","Triage, SLA, assignment, evidence and verification"),("3 min","Asset & PPM","Asset passport, calendar and compliance"),("3 min","Performance & reports","KPI drill-down and landlord output"),("2 min","Close","Return to Companion and recap proof chain")]
    y=2.45
    for i,(tm,stage,msg) in enumerate(rows):
        fill=WHITE if i%2==0 else RGBColor(234,241,247); rect(s,0.72,y,11.9,0.48,fill,fill,False)
        add_text(s,tm,0.9,y+0.1,0.75,0.22,11,RED,True); add_text(s,stage,1.75,y+0.1,2.2,0.22,11,INK,True); add_text(s,msg,4.05,y+0.1,8.2,0.22,11,MUTED)
        y+=0.5
    proof_bar(s,"Perform one meaningful action live; use prepared records for the broader lifecycle and reporting story.")

    s = base_slide(prs,"Presenter guardrails","What to claim—and what to position as roadmap","Credibility matters more than covering every feature.",14)
    card(s,0.7,2.48,3.78,3.45,"Show live","Property scoping, tenant/lease records, issue reporting, work orders, PPM, service profiles/providers, reports and Companion workflows.",GREEN,"✓")
    card(s,4.77,2.48,3.78,3.45,"Describe carefully","External email/SMS delivery, live payment gateways and scheduled automations only when the environment is configured and tested.",AMBER,"!")
    card(s,8.84,2.48,3.78,3.45,"Position as integration","Specific smart-meter devices, third-party accounting systems and provider-specific APIs require vendor credentials and interface work.",RED,"→")
    proof_bar(s,"Never claim a feature is live unless it works in the environment being demonstrated.")

    s = base_slide(prs,"Demo readiness","Final pre-flight checklist","Run this on the same device, browser and network that will be used in the meeting.",15)
    checks=["Login and session persistence","Correct organization and selected property","Property, tenant, asset and provider records populated","Open, due and overdue work represented","Tenant issue submission and generated reference","Work-order detail and SLA fields","PPM and compliance calendar","Companion tasks and notifications","PDF report opens and is legible","Backup screenshots/video available"]
    for i,c in enumerate(checks):
        x=0.75+(i%2)*6.05; y=2.4+(i//2)*0.72
        rect(s,x,y,5.72,0.55,WHITE,LINE); add_text(s,"□",x+0.16,y+0.12,0.3,0.25,17,RED,True); add_text(s,c,x+0.52,y+0.11,4.95,0.25,12,INK,True)
    proof_bar(s,"Stop and switch to the backup path if latency, authentication or integrations are unstable.")

    s = base_slide(prs,"Close","CeEst is an operating control system","A single live record connects tenants, staff, providers, assets, money and management evidence.",16)
    rect(s,0.7,2.55,11.95,1.05,NAVY,NAVY)
    add_text(s,"RECORD  →  PLAN  →  ASSIGN  →  EXECUTE  →  VERIFY  →  MEASURE  →  REPORT  →  IMPROVE",0.92,2.88,11.5,0.35,16,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,"The facility is onboarded. Work is planned. Issues are controlled. Performance is measured. Reports are ready.",1.15,4.25,11.0,0.85,24,INK,True,PP_ALIGN.CENTER)
    add_text(s,"Questions & live scenarios",4.55,5.65,4.2,0.42,17,RED,True,PP_ALIGN.CENTER)
    proof_bar(s,"One operating record. Complete estate control.")

    OUT.mkdir(parents=True, exist_ok=True)
    path=OUT / "CeEst_ERP_Live_Demo_Presentation.pptx"
    prs.save(path)
    return path


if __name__ == "__main__":
    print(build())
